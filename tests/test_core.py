import base64
import subprocess
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import llm, store
from app.analysis import anomalies, month_stats
from app.parser import InvoiceParseError, extract_xml, parse_fatturapa

SAMPLES = Path(__file__).resolve().parent.parent / "samples"
TODAY = date(2026, 10, 17)


def sample(name: str) -> bytes:
    return next(SAMPLES.glob(f"{name}-*.xml")).read_bytes()


def test_parses_supplier_amounts_and_payments():
    (inv,) = parse_fatturapa(sample("01"))
    assert inv.supplier.name == "Nuvola Hosting S.r.l."
    assert inv.supplier.vat_number == "10000000001"
    assert inv.customer.name == "Giulia Rossi"
    assert inv.number == "2026/0412"
    assert [l.description for l in inv.lines] == ["Hosting VPS annuale", "Dominio .it rinnovo"]
    assert float(inv.taxable) == 252.00
    assert float(inv.tax) == 55.44
    assert float(inv.total) == 307.44
    assert inv.payments[0].due_date == "2026-10-31"


def test_multiple_bodies_in_one_file():
    xml = sample("01").decode()
    start = xml.index("<FatturaElettronicaBody>")
    end = xml.index("</FatturaElettronicaBody>") + len("</FatturaElettronicaBody>")
    body = xml[start:end]
    doubled = xml.replace(body, body + body.replace("2026/0412", "2026/0413"))
    assert [i.number for i in parse_fatturapa(doubled.encode())] == ["2026/0412", "2026/0413"]


def test_rejects_non_invoice_and_entity_bombs():
    with pytest.raises(InvoiceParseError):
        parse_fatturapa(b"<root/>")
    bomb = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><FatturaElettronica>&a;</FatturaElettronica>'
    with pytest.raises(InvoiceParseError):
        parse_fatturapa(bomb)


@pytest.fixture
def p7m(tmp_path):
    key, cert = tmp_path / "k.pem", tmp_path / "c.pem"
    subprocess.run(
        ["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-subj", "/CN=test",
         "-keyout", key, "-out", cert, "-days", "1"],
        check=True, capture_output=True,
    )
    src = tmp_path / "f.xml"
    src.write_bytes(sample("03"))
    out = tmp_path / "f.xml.p7m"
    subprocess.run(
        ["openssl", "cms", "-sign", "-binary", "-nodetach", "-outform", "DER",
         "-in", src, "-signer", cert, "-inkey", key, "-out", out],
        check=True, capture_output=True,
    )
    return out.read_bytes()


def test_extracts_signed_p7m_der_and_base64(p7m):
    assert parse_fatturapa(p7m)[0].supplier.name == "Studio Bianchi Commercialisti"
    assert b"FatturaElettronica" in extract_xml(base64.encodebytes(p7m))


def _record(i, name):
    return {"id": i, "invoice": parse_fatturapa(sample(name))[0].to_dict(), "paid": False, "category": "x"}


def test_anomalies_duplicate_vat_overdue():
    records = [_record(1, "01"), _record(8, "08"), _record(7, "07"), _record(9, "09")]
    kinds = lambda r: {a["kind"] for a in anomalies(r, records, TODAY)}
    assert "duplicate" in kinds(records[1])
    assert "vat_mismatch" in kinds(records[2])
    assert "overdue" in kinds(records[3])
    assert kinds(records[0]) == {"duplicate"}


def test_month_stats_upcoming_and_totals():
    records = [_record(i, f"{i:02d}") for i in range(1, 11)]
    stats = month_stats(records, "2026-10", TODAY)
    assert stats["invoice_count"] == 9
    assert [p["due_date"] for p in stats["upcoming_payments"]][:2] == ["2026-10-20", "2026-10-31"]


def test_rules_fallback_categorizes():
    inv = parse_fatturapa(sample("02"))[0].to_dict()
    assert llm.categorize_by_rules(inv)["category"] == "travel"


def test_api_upload_and_export(tmp_path, monkeypatch):
    monkeypatch.setenv("INVOICELENS_DB", str(tmp_path / "t.db"))
    monkeypatch.setenv("LLM_BASE_URL", "http://127.0.0.1:9")  # unreachable: forces rules fallback
    from app.main import app

    client = TestClient(app)
    res = client.post("/api/invoices", files=[("files", ("01.xml", sample("01"), "text/xml"))])
    assert res.status_code == 200
    assert res.json()["stored"][0]["category_source"] == "rules"
    assert client.get("/api/invoices").json()[0]["invoice"]["number"] == "2026/0412"
    assert "Nuvola Hosting" in client.get("/api/export.csv").text
    bad = client.post("/api/invoices", files=[("files", ("x.xml", b"<nope/>", "text/xml"))])
    assert bad.json()["errors"]


class _FakeResponse:
    def __init__(self, content):
        self._content = content

    def raise_for_status(self):
        pass

    def json(self):
        return {
            "choices": [{"message": {"content": self._content}}],
            "usage": {"prompt_tokens": 10, "completion_tokens": 5},
        }


def test_model_path_sends_openai_compatible_request(monkeypatch):
    sent = {}

    def fake_post(url, json, headers, timeout):
        sent.update(url=url, body=json)
        return _FakeResponse('Sure! {"category": "software", "confidence": 0.9, "reason": "hosting"}')

    monkeypatch.setattr(llm.httpx, "post", fake_post)
    monkeypatch.setenv("LLM_MODEL", "gemma3:4b")
    inv = parse_fatturapa(sample("01"))[0].to_dict()
    result = llm.categorize(inv)
    assert result == {"category": "software", "confidence": 0.9, "reason": "hosting", "source": "gemma3:4b"}
    assert sent["url"].endswith("/v1/chat/completions")
    assert sent["body"]["model"] == "gemma3:4b"
    assert "Hosting VPS annuale" in sent["body"]["messages"][1]["content"]


def test_model_unknown_category_becomes_other(monkeypatch):
    monkeypatch.setattr(llm.httpx, "post", lambda *a, **k: _FakeResponse('{"category": "pizza"}'))
    inv = parse_fatturapa(sample("01"))[0].to_dict()
    assert llm.categorize(inv)["category"] == "other"


def _client(tmp_path, monkeypatch):
    monkeypatch.setenv("INVOICELENS_DB", str(tmp_path / "t.db"))
    monkeypatch.setenv("LLM_BASE_URL", "http://127.0.0.1:9")
    from app.main import app

    return TestClient(app)


def test_zip_bulk_import_skips_receipts_and_reimports(tmp_path, monkeypatch):
    import io
    import zipfile

    client = _client(tmp_path, monkeypatch)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ottobre/IT10000000001_a1.xml", sample("01"))
        z.writestr("ottobre/IT10000000002_a2.xml", sample("02"))
        z.writestr("ottobre/IT10000000001_a1_MT_001.xml", b"<MetadatiInvioFile/>")
        z.writestr("__MACOSX/._x.xml", b"junk")
        z.writestr("notes.txt", b"ignored")
    files = [("files", ("ottobre.zip", buf.getvalue(), "application/zip"))]
    first = client.post("/api/invoices", files=files).json()
    assert len(first["stored"]) == 2 and first["errors"] == []
    again = client.post("/api/invoices", files=files).json()
    assert again["stored"] == []
    assert len(client.get("/api/invoices").json()) == 2


def test_bad_zip_reported(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    res = client.post("/api/invoices", files=[("files", ("x.zip", b"nope", "application/zip"))]).json()
    assert res["errors"][0]["error"] == "not a valid ZIP file"


def test_deadlines_calendar_excludes_paid(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    client.post("/api/invoices", files=[("files", ("a.xml", sample("01"), "text/xml")),
                                        ("files", ("b.xml", sample("03"), "text/xml"))])
    ids = {r["invoice"]["number"]: r["id"] for r in client.get("/api/invoices").json()}
    client.post(f"/api/invoices/{ids['57']}/paid", json={"paid": True})
    body = client.get("/api/deadlines.ics").text
    assert "DTSTART;VALUE=DATE:20261031" in body
    assert "Studio Bianchi" not in body
    assert "Nuvola Hosting S.r.l. 307.44 EUR" in body


def test_patch_overrides_category_and_paid(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    client.post("/api/invoices", files=[("files", ("a.xml", sample("05"), "text/xml"))])
    (rec,) = client.get("/api/invoices").json()
    assert rec["category_source"] == "rules"
    res = client.patch(f"/api/invoices/{rec['id']}", json={"category": "office"})
    assert res.status_code == 200
    body = res.json()
    assert body["category"] == "office" and body["category_source"] == "manual"
    assert body["paid"] is False
    assert any(a["kind"] == "overdue" for a in body["anomalies"])
    paid = client.patch(f"/api/invoices/{rec['id']}", json={"paid": True}).json()
    assert paid["paid"] is True and paid["category"] == "office"
    assert not any(a["kind"] == "overdue" for a in paid["anomalies"])
    assert client.patch(f"/api/invoices/{rec['id']}", json={"category": "pizza"}).status_code == 400
    assert client.patch(f"/api/invoices/{rec['id']}", json={}).status_code == 400
    assert client.patch("/api/invoices/999", json={"paid": True}).status_code == 404


def test_upload_reports_skipped_duplicates(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    files = [("files", ("a.xml", sample("01"), "text/xml"))]
    assert client.post("/api/invoices", files=files).json()["skipped"] == []
    again = client.post("/api/invoices", files=files).json()
    assert again["stored"] == [] and again["errors"] == []
    assert again["skipped"] == [
        {"file": "a.xml", "number": "2026/0412", "supplier": "Nuvola Hosting S.r.l.", "reason": "already imported"}
    ]


def test_delete_invoice_and_demo_is_read_only(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    client.post("/api/invoices", files=[("files", ("a.xml", sample("01"), "text/xml"))])
    (rec,) = client.get("/api/invoices").json()
    assert client.delete(f"/api/invoices/{rec['id']}").json() == {"ok": True}
    assert client.get("/api/invoices").json() == []
    assert client.delete(f"/api/invoices/{rec['id']}").status_code == 404
    monkeypatch.setenv("INVOICELENS_DEMO", "1")
    assert client.delete("/api/invoices/1").status_code == 403


def test_config_reports_model_status_and_samples(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    llm._status_cache.clear()
    cfg = client.get("/api/config").json()
    assert cfg["model"] == llm.model_name()
    assert cfg["model_reachable"] is False
    assert cfg["uploads_enabled"] is True and cfg["can_load_samples"] is True
    assert "software" in cfg["categories"]

    class _Ok:
        status_code = 200

    llm._status_cache.clear()
    monkeypatch.setattr(llm.httpx, "get", lambda *a, **k: _Ok())
    assert client.get("/api/config").json()["model_reachable"] is True
    llm._status_cache.clear()


def test_samples_load_only_into_empty_local_store(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    res = client.post("/api/demo/load-samples")
    assert res.status_code == 200 and res.json()["loaded"] == 10
    assert client.get("/api/config").json()["can_load_samples"] is False
    assert client.post("/api/demo/load-samples").status_code == 403
    monkeypatch.setenv("INVOICELENS_DEMO", "1")
    assert client.post("/api/demo/load-samples").json()["loaded"] == 10
    assert len(client.get("/api/invoices").json()) == 10


def test_retries_once_on_timeout(monkeypatch):
    calls = []

    def flaky_post(url, json, headers, timeout):
        calls.append(1)
        if len(calls) == 1:
            raise llm.httpx.ReadTimeout("model still loading")
        return _FakeResponse('{"category": "software", "confidence": 0.9, "reason": "hosting"}')

    monkeypatch.setattr(llm.httpx, "post", flaky_post)
    inv = parse_fatturapa(sample("01"))[0].to_dict()
    assert llm.categorize(inv)["category"] == "software"
    assert len(calls) == 2


def test_summary_facts_are_plain_and_only_list_real_overdue():
    records = [_record(i, f"{i:02d}") for i in range(1, 11)]
    facts = llm.summary_facts(month_stats(records, "2026-10", TODAY))
    overdue = next(l for l in facts.splitlines() if l.startswith("Overdue:"))
    assert "Coworking Dorsoduro" in overdue and "Accademia Digitale" not in overdue
    assert "VAT that can be reclaimed" in facts and "{" not in facts


def test_accountant_pack_has_summary_csv_checks_and_calendar(tmp_path, monkeypatch):
    import io
    import zipfile

    client = _client(tmp_path, monkeypatch)
    client.post("/api/invoices", files=[("files", (f"{n}.xml", sample(n), "text/xml")) for n in ("01", "07", "08", "09")])
    res = client.get("/api/accountant-pack.zip?month=2026-10")
    assert res.status_code == 200
    assert "invoicelens-commercialista-2026-10.zip" in res.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(res.content))
    assert sorted(z.namelist()) == ["da-verificare-2026-10.csv", "fatture-2026-10.csv", "fatture-2026-10.xlsx",
                                    "riepilogo-2026-10.html", "scadenze.ics"]

    invoices = z.read("fatture-2026-10.csv").decode("utf-8-sig").splitlines()
    assert len(invoices) == 1 + 3  # header + the three October invoices; Energia Chiara is September
    assert "Accademia Digitale" in "\n".join(invoices) and "Energia Chiara" not in "\n".join(invoices)
    assert "199,00;40,00;239,00" in "\n".join(invoices)  # Italian decimal comma, amounts from the XML

    checks = z.read("da-verificare-2026-10.csv").decode("utf-8-sig")
    assert "IVA al 22% su 199,00 dovrebbe essere 43,78, la fattura riporta 40,00" in checks
    assert checks.count("Possibile doppione") == 2
    assert "imponibile × aliquota, arrotondato al centesimo" in checks  # the rule travels with the flag

    summary = z.read("riepilogo-2026-10.html").decode()
    assert "Fatture passive di ottobre 2026" in summary
    assert "Solo indicativo" in summary and "forfettario" in summary
    assert "BEGIN:VCALENDAR" in z.read("scadenze.ics").decode()
    assert client.get("/api/accountant-pack.zip?month=10-2026").status_code == 400


def test_invoices_carry_plain_explanation_and_tax_hint(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    client.post("/api/invoices", files=[("files", ("07.xml", sample("07"), "text/xml"))])
    (r,) = client.get("/api/invoices").json()
    assert r["explanation"].startswith("Accademia Digitale S.r.l. billed you €199.00 plus €40.00 VAT (22%)")
    assert "by bank transfer" in r["explanation"]
    assert "to check before you pay" in r["explanation"]
    vat = next(a for a in r["anomalies"] if a["kind"] == "vat_mismatch")
    assert "taxable × rate" in vat["rule"] and "€0.02" in vat["rule"]
    patched = client.patch(f"/api/invoices/{r['id']}", json={"category": "telecom"}).json()
    assert patched["tax_hint"].startswith("Phone and internet: 80% deductible")
    assert "forfettario" in client.get("/api/config").json()["tax_hint_disclaimer"]


def test_credit_note_is_explained_as_money_back():
    from app.accountant import explain

    inv = parse_fatturapa(sample("03"))[0].to_dict()
    inv["doc_type"] = "TD04"
    text = explain({"invoice": inv, "paid": True}, [])
    assert "credit note" in text and "not a new cost" in text and "marked it as paid" in text


def test_csv_export_neutralizes_spreadsheet_formulas(tmp_path, monkeypatch):
    from app.accountant import csv_safe

    assert csv_safe("=HYPERLINK(\"http://x\")") == "'=HYPERLINK(\"http://x\")"
    assert csv_safe("+39 caffè") == "'+39 caffè" and csv_safe("Nuvola S.r.l.") == "Nuvola S.r.l."
    client = _client(tmp_path, monkeypatch)
    evil = sample("01").replace(b"Nuvola Hosting S.r.l.", b"=cmd|' /C calc'!A0")
    client.post("/api/invoices", files=[("files", ("evil.xml", evil, "text/xml"))])
    assert ";'=cmd|" in client.get("/api/export.csv").text


def test_excel_report_has_summary_invoices_and_checks_with_rules(tmp_path, monkeypatch):
    import io

    from openpyxl import load_workbook

    client = _client(tmp_path, monkeypatch)
    evil = sample("08").replace(b"Nuvola Hosting S.r.l.", b"=1+1 Hosting")
    client.post("/api/invoices", files=[("files", ("01.xml", sample("01"), "text/xml")),
                                        ("files", ("07.xml", sample("07"), "text/xml")),
                                        ("files", ("evil.xml", evil, "text/xml"))])
    res = client.get("/api/report.xlsx?month=2026-10")
    assert res.status_code == 200 and "invoicelens-2026-10.xlsx" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Riepilogo", "Fatture", "Da verificare"]
    summary = [c.value for row in wb["Riepilogo"].iter_rows() for c in row if c.value is not None]
    assert "Fatture passive di ottobre 2026" in summary and 3 in summary
    rows = list(wb["Fatture"].iter_rows(values_only=True))
    accademia = next(r for r in rows if r[3] == "Accademia Digitale S.r.l.")
    assert accademia[7:10] == (199.0, 40.0, 239.0)  # real numbers, not text
    assert any(r[3] == "=1+1 Hosting" for r in rows)  # kept as text, not evaluated as a formula
    assert wb["Fatture"]["D4"].data_type == "s"
    checks = list(wb["Da verificare"].iter_rows(values_only=True))
    assert any("43,78" in (r[4] or "") and "0,02" in (r[5] or "") for r in checks)


def test_wipe_deletes_everything_only_with_confirmation(tmp_path, monkeypatch):
    client = _client(tmp_path, monkeypatch)
    client.post("/api/invoices", files=[("files", ("01.xml", sample("01"), "text/xml"))])
    assert client.post("/api/wipe", json={"confirm": "yes"}).status_code == 400
    assert len(client.get("/api/invoices").json()) == 1
    res = client.post("/api/wipe", json={"confirm": "DELETE"}).json()
    assert res["deleted"] == 1
    assert client.get("/api/invoices").json() == []
    assert b"Nuvola" not in (tmp_path / "t.db").read_bytes()  # vacuumed, not just unlinked rows
    monkeypatch.setenv("INVOICELENS_DEMO", "1")
    assert client.post("/api/wipe", json={"confirm": "DELETE"}).status_code == 403
