"""Everything the accountant needs for one month, in one ZIP; plus plain-language explanations and
indicative deductibility hints. All of it is built in code from the parsed invoices: no model call."""

from __future__ import annotations

import csv
import html
import io
import zipfile
from collections import defaultdict
from datetime import date

from .analysis import anomalies, month_stats

DISCLAIMER_EN = (
    "Indicative only: general rules for a self-employed professional in the ordinary or simplified regime. "
    "In the regime forfettario costs are not deducted and VAT is not reclaimed. Your accountant decides."
)
DISCLAIMER_IT = (
    "Solo indicativo: regole generali per un professionista in regime ordinario o semplificato. "
    "Nel regime forfettario i costi non si deducono e l'IVA non si detrae. Decide il commercialista."
)

# (English for the app, Italian for the accountant pack). Deliberately short and hedged.
TAX_HINTS = {
    "software": (
        "Usually 100% deductible when used for your work; VAT reclaimable.",
        "Di norma deducibile al 100% se inerente all'attività; IVA detraibile.",
    ),
    "hardware": (
        "A work tool: deducted over several years (depreciation), or in full this year if it costs "
        "up to €516.46. VAT reclaimable.",
        "Bene strumentale: si deduce con l'ammortamento, oppure tutto nell'anno se costa fino a 516,46 €. "
        "IVA detraibile.",
    ),
    "travel": (
        "Work trips are usually deductible. Hotels follow the 75% rule (see meals); car costs such as fuel "
        "and tolls are 20% deductible with 40% of the VAT. From 2025, hotels and taxis in Italy count only "
        "if paid by card or bank transfer.",
        "Trasferte di lavoro di norma deducibili. Alberghi: regola del 75% (vedi vitto); spese auto "
        "(carburante, pedaggi) deducibili al 20% con IVA detraibile al 40%. Dal 2025 alberghi e taxi in "
        "Italia solo se pagati con mezzi tracciabili.",
    ),
    "meals": (
        "Restaurants and hotels: 75% deductible, capped in total at 2% of the fees you received in the year. "
        "VAT reclaimable with an invoice. From 2025, in Italy only if paid by card or bank transfer.",
        "Vitto e alloggio: deducibili al 75%, entro il 2% dei compensi dell'anno. IVA detraibile con fattura. "
        "Dal 2025, in Italia, solo se pagati con mezzi tracciabili.",
    ),
    "office": (
        "Stationery and coworking desks: usually 100% deductible; VAT reclaimable.",
        "Cancelleria e postazioni coworking: di norma deducibili al 100%; IVA detraibile.",
    ),
    "utilities": (
        "Office at home and no other office in the same town: 50%. A separate office: 100%.",
        "Studio in casa senza altro studio nello stesso comune: 50%. Studio separato: 100%.",
    ),
    "telecom": (
        "Phone and internet: 80% deductible; 50% of the VAT is reclaimable.",
        "Telefonia e internet: deducibili all'80%; IVA detraibile al 50%.",
    ),
    "professional_services": (
        "Accountant, lawyer, consultants: usually 100% deductible; VAT reclaimable.",
        "Commercialista, avvocato, consulenti: di norma deducibili al 100%; IVA detraibile.",
    ),
    "training": (
        "Courses and conferences, including travel and stay for them: 100% deductible up to €10,000 a year.",
        "Corsi e convegni, compresi viaggio e soggiorno: deducibili al 100% fino a 10.000 € l'anno.",
    ),
    "marketing": (
        "Advertising: usually 100% deductible. Client gifts and hospitality are 'spese di rappresentanza' "
        "with their own limits.",
        "Pubblicità: di norma deducibile al 100%. Omaggi e ospitalità sono spese di rappresentanza, con "
        "limiti propri.",
    ),
    "rent": (
        "Office rent: 100%. A home that is also your office (no other office in the same town): 50%.",
        "Affitto dello studio: 100%. Casa usata anche come studio (senza altro studio nel comune): 50%.",
    ),
    "other": (
        "No general rule for this category: ask your accountant.",
        "Nessuna regola generale per questa categoria: da valutare con il commercialista.",
    ),
}

CATEGORY_IT = {
    "software": "Software e cloud",
    "hardware": "Hardware",
    "travel": "Viaggi e trasporti",
    "meals": "Vitto e alloggio",
    "office": "Ufficio e coworking",
    "utilities": "Utenze",
    "telecom": "Telefonia e internet",
    "professional_services": "Servizi professionali",
    "training": "Formazione",
    "marketing": "Pubblicità e marketing",
    "rent": "Affitto",
    "other": "Altro",
}

ANOMALY_IT = {
    "duplicate": "Possibile doppione",
    "amount_conflict": "Stesso numero, importi diversi",
    "possible_duplicate": "Stesso importo a pochi giorni",
    "vat_mismatch": "IVA non coerente con l'aliquota",
    "total_mismatch": "Totale diverso da imponibile + IVA",
    "overdue": "Pagamento scaduto",
}

DOC_TYPES = {
    "TD01": "invoice",
    "TD02": "advance invoice",
    "TD04": "credit note",
    "TD05": "debit note",
    "TD06": "fee note (parcella)",
    "TD24": "deferred invoice",
    "TD25": "deferred invoice",
    "TD26": "invoice for the sale of a business asset",
    "PDF": "PDF invoice",
}

PAYMENT_METHODS = {
    "MP01": "cash",
    "MP02": "cheque",
    "MP05": "bank transfer",
    "MP08": "card",
    "MP12": "RIBA",
    "MP19": "SEPA direct debit",
    "MP20": "SEPA direct debit",
    "MP21": "SEPA direct debit",
    "MP23": "PagoPA",
}

MONTHS_IT = ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto",
             "settembre", "ottobre", "novembre", "dicembre"]


def tax_hint(category: str) -> str:
    return TAX_HINTS.get(category, TAX_HINTS["other"])[0]


def csv_safe(value) -> str:
    """Keep a supplier-controlled string from being run as a spreadsheet formula (CSV injection)."""
    text = str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _eur(v: float) -> str:
    return f"€{v:,.2f}"


def _fmt_date(iso: str) -> str:
    try:
        d = date.fromisoformat(iso)
        return f"{d.day} {d.strftime('%b %Y')}"
    except ValueError:
        return iso or "no date"


def explain(record: dict, found: list[dict]) -> str:
    """The invoice in two or three plain sentences, for someone who does not read FatturaPA."""
    inv = record["invoice"]
    kind = DOC_TYPES.get(inv.get("doc_type", ""), "invoice")
    lines = [l["description"] for l in inv["lines"] if l["description"]]
    what = lines[0] if lines else "unspecified goods or services"
    if len(lines) > 1:
        what += f" and {len(lines) - 1} more line{'s' if len(lines) > 2 else ''}"
    rates = sorted({f"{v['rate']:g}%" for v in inv["vat_summary"]})
    vat = f"{_eur(inv['tax'])} VAT ({', '.join(rates)})" if rates else f"{_eur(inv['tax'])} VAT"
    supplier = inv["supplier"]["name"] or "The supplier"
    if kind == "credit note":
        out = [f"{supplier} sent you a credit note: {_eur(inv['taxable'])} plus {vat} for {what}. "
               "It reduces what you owe or gives money back; it is not a new cost."]
    else:
        out = [f"{supplier} billed you {_eur(inv['taxable'])} plus {vat} for {what}, "
               f"{_eur(inv['total'])} in total ({kind}, n. {inv['number']}, {_fmt_date(inv['date'])})."]

    if record.get("paid"):
        out.append("You marked it as paid.")
    else:
        dues = [p for p in inv["payments"] if p["due_date"]]
        if dues:
            p = dues[0]
            how = PAYMENT_METHODS.get(p.get("method", ""), "")
            out.append(f"Pay {_eur(p['amount'])} by {_fmt_date(p['due_date'])}" + (f" by {how}." if how else "."))
        else:
            out.append("The invoice gives no due date.")

    if found:
        n = len(found)
        out.append(f"InvoiceLens found {n} thing{'s' if n > 1 else ''} to check before you pay or deduct it.")
    else:
        out.append("The arithmetic checks passed.")
    return " ".join(out)


def _month_label(month: str) -> str:
    try:
        y, m = month.split("-")
        return f"{MONTHS_IT[int(m) - 1]} {y}"
    except (ValueError, IndexError):
        return month


def _num(v: float) -> str:
    return f"{v:.2f}".replace(".", ",")


def _num_it(v: float) -> str:
    """1873.11 -> 1.873,11 for people; the CSV keeps 1873,11 so spreadsheets parse it."""
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    for row in rows:
        writer.writerow([csv_safe(c) for c in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel opens the accents correctly


def _msg(a: dict) -> str:
    return a.get("message_it") or a["message"]


def _source_it(src: str) -> str:
    return {"manual": "manuale", "rules": "regole"}.get(src, src)


def _due(inv: dict) -> str:
    return next((p["due_date"] for p in inv["payments"] if p["due_date"]), "")


def accountant_pack(records: list[dict], month: str, ics: str, today: date | None = None) -> bytes:
    """One ZIP for the commercialista: a printable summary, the month's invoices as CSV, the items to check,
    and the unpaid due dates as a calendar file."""
    today = today or date.today()
    stats = month_stats(records, month, today)
    in_month = [r for r in records if r["invoice"]["date"].startswith(month)]
    flagged = {r["id"]: anomalies(r, records, today) for r in in_month}

    invoices_csv = [["id", "data", "numero", "tipo", "fornitore", "partita_iva", "categoria", "fonte_categoria",
                     "imponibile", "iva", "totale", "scadenza", "pagata", "da_verificare", "nota_deducibilita"]]
    for r in sorted(in_month, key=lambda r: (r["invoice"]["date"], r["id"])):
        inv = r["invoice"]
        invoices_csv.append([
            r["id"], inv["date"], inv["number"], inv.get("doc_type", ""), inv["supplier"]["name"],
            inv["supplier"]["vat_number"], CATEGORY_IT.get(r["category"], r["category"]), _source_it(r["category_source"]),
            _num(inv["taxable"]), _num(inv["tax"]), _num(inv["total"]), _due(inv), "sì" if r["paid"] else "no",
            " | ".join(f"{ANOMALY_IT.get(a['kind'], a['kind'])}: {_msg(a)}" for a in flagged[r["id"]]),
            TAX_HINTS.get(r["category"], TAX_HINTS["other"])[1],
        ])

    checks_csv = [["id", "fornitore", "numero", "data", "problema", "dettaglio", "regola"]]
    for r in in_month:
        for a in flagged[r["id"]]:
            checks_csv.append([r["id"], r["invoice"]["supplier"]["name"], r["invoice"]["number"],
                               r["invoice"]["date"], ANOMALY_IT.get(a["kind"], a["kind"]), _msg(a), a.get("rule_it", "")])

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"riepilogo-{month}.html", _summary_html(stats, in_month, flagged, today))
        z.writestr(f"fatture-{month}.csv", _csv(invoices_csv))
        z.writestr(f"da-verificare-{month}.csv", _csv(checks_csv))
        z.writestr(f"fatture-{month}.xlsx", report_xlsx(records, month, today))
        z.writestr("scadenze.ics", ics)
    return out.getvalue()


def _summary_html(stats: dict, in_month: list[dict], flagged: dict, today: date) -> str:
    e = html.escape
    by_cat: dict[str, list[float]] = defaultdict(lambda: [0, 0.0, 0.0])
    sources: dict[str, int] = defaultdict(int)
    for r in in_month:
        c = by_cat[r["category"]]
        c[0] += 1
        c[1] += r["invoice"]["taxable"]
        c[2] += r["invoice"]["tax"]
        src = r["category_source"]
        sources[_source_it(src)] += 1
    cat_rows = "".join(
        f"<tr><td>{e(CATEGORY_IT.get(k, k))}</td><td class=n>{v[0]}</td><td class=n>{_num_it(v[1])}</td>"
        f"<td class=n>{_num_it(v[2])}</td><td class=h>{e(TAX_HINTS.get(k, TAX_HINTS['other'])[1])}</td></tr>"
        for k, v in sorted(by_cat.items(), key=lambda kv: -kv[1][1])
    )
    checks = "".join(
        f"<li><b>{e(r['invoice']['supplier']['name'])}</b> n. {e(r['invoice']['number'])}: "
        f"{e(ANOMALY_IT.get(a['kind'], a['kind']))}. <span class=m>{e(_msg(a))}</span><br><small class=m>Regola: {e(a.get('rule_it', ''))}</small></li>"
        for r in in_month for a in flagged[r["id"]]
    ) or "<li>Nessun problema trovato dai controlli.</li>"
    inv_rows = "".join(
        f"<tr><td>{e(r['invoice']['date'])}</td><td>{e(r['invoice']['supplier']['name'])}</td>"
        f"<td>{e(r['invoice']['number'])}</td><td>{e(CATEGORY_IT.get(r['category'], r['category']))}</td>"
        f"<td class=n>{_num_it(r['invoice']['taxable'])}</td><td class=n>{_num_it(r['invoice']['tax'])}</td>"
        f"<td class=n>{_num_it(r['invoice']['total'])}</td><td>{'sì' if r['paid'] else 'no'}</td>"
        f"<td>{'⚠ ' + str(len(flagged[r['id']])) if flagged[r['id']] else ''}</td></tr>"
        for r in sorted(in_month, key=lambda r: (r["invoice"]["date"], r["id"]))
    )
    due = "".join(
        f"<li>{e(p['due_date'])}: {e(p['supplier'])}, {_num_it(p['amount'])} €</li>" for p in stats["upcoming_payments"]
    ) or "<li>Nessuna scadenza nei prossimi 30 giorni.</li>"
    src_text = ", ".join(f"{n} da {e(s)}" for s, n in sorted(sources.items())) or "nessuna"
    label = _month_label(stats["month"])
    total = round(stats["taxable_total"] + stats["vat_total"], 2)
    return f"""<!doctype html>
<html lang="it"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Riepilogo fatture passive · {e(label)}</title>
<style>
body{{font:14px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif;color:#1b1f24;background:#fff;max-width:960px;margin:24px auto;padding:0 16px}}
h1{{font-size:22px;margin:0}} h2{{font-size:15px;margin:22px 0 6px;border-bottom:1px solid #d9dee4;padding-bottom:4px}}
.sub,.m{{color:#5b6570}} .kpi{{display:flex;gap:10px;flex-wrap:wrap;margin:14px 0}}
.kpi div{{border:1px solid #d9dee4;border-radius:8px;padding:8px 12px;min-width:140px}} .kpi b{{display:block;font-size:18px}}
table{{border-collapse:collapse;width:100%;font-size:13px}} th,td{{text-align:left;padding:5px 6px;border-bottom:1px solid #eceff2;vertical-align:top}}
th{{color:#5b6570;font-weight:600}} .n{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}} .h{{font-size:12px;color:#3d4650}}
.note{{background:#f6f7f9;border-radius:8px;padding:8px 12px;font-size:12.5px;color:#3d4650}}
@media print{{body{{margin:0}} .noprint{{display:none}}}}
</style></head><body>
<h1>Fatture passive di {e(label)}</h1>
<div class=sub>Preparato con InvoiceLens il {today.strftime('%d/%m/%Y')}. Importi in euro.</div>
<p class="note noprint">Per un PDF: stampa questa pagina dal browser (Ctrl/Cmd+P → Salva come PDF).</p>
<div class=kpi><div>Fatture<b>{stats['invoice_count']}</b></div><div>Imponibile<b>{_num_it(stats['taxable_total'])}</b></div>
<div>IVA<b>{_num_it(stats['vat_total'])}</b></div><div>Totale<b>{_num_it(total)}</b></div><div>Da verificare<b>{len(stats['anomalies'])}</b></div></div>
<h2>Da verificare</h2><ul>{checks}</ul>
<p class=note>Controlli fatti con l'aritmetica sui dati XML, non da un modello: doppioni (stesso fornitore, numero e anno), IVA diversa da imponibile × aliquota (tolleranza 2 centesimi), totale diverso da imponibile + IVA (spesso ritenuta o bollo), pagamenti scaduti.</p>
<h2>Per categoria</h2>
<table><thead><tr><th>Categoria</th><th class=n>Fatture</th><th class=n>Imponibile</th><th class=n>IVA</th><th>Nota sulla deducibilità (indicativa)</th></tr></thead><tbody>{cat_rows}</tbody></table>
<p class=note>{e(DISCLAIMER_IT)} Categorie: {src_text}. Sono proposte, da confermare.</p>
<h2>Scadenze nei prossimi 30 giorni</h2><ul>{due}</ul>
<h2>Elenco fatture</h2>
<table><thead><tr><th>Data</th><th>Fornitore</th><th>Numero</th><th>Categoria</th><th class=n>Imponibile</th><th class=n>IVA</th><th class=n>Totale</th><th>Pagata</th><th></th></tr></thead><tbody>{inv_rows}</tbody></table>
<p class=sub>Nel pacchetto: fatture-{e(stats['month'])}.xlsx (Excel, tre fogli), fatture-{e(stats['month'])}.csv (separatore ; e virgola decimale), da-verificare-{e(stats['month'])}.csv, scadenze.ics.</p>
</body></html>
"""


def report_xlsx(records: list[dict], month: str, today: date | None = None) -> bytes:
    """The month as an Excel workbook: summary by category, every invoice, and what to check (with the rule)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    today = today or date.today()
    stats = month_stats(records, month, today)
    in_month = sorted((r for r in records if r["invoice"]["date"].startswith(month)),
                      key=lambda r: (r["invoice"]["date"], r["id"]))
    flagged = {r["id"]: anomalies(r, records, today) for r in in_month}
    money = "#,##0.00"
    bold = Font(bold=True)

    def put(ws, row: list, header: bool = False) -> None:
        ws.append([None] * len(row))
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # never let a supplier name become a formula
            elif isinstance(value, float):
                cell.number_format = money
            if header:
                cell.font = bold

    wb = Workbook()
    ws = wb.active
    ws.title = "Riepilogo"
    put(ws, [f"Fatture passive di {_month_label(month)}"], header=True)
    put(ws, [f"Preparato con InvoiceLens il {today.strftime('%d/%m/%Y')}. Importi in euro."])
    put(ws, [])
    for label, value in (("Fatture", stats["invoice_count"]), ("Imponibile", float(stats["taxable_total"])),
                         ("IVA", float(stats["vat_total"])),
                         ("Totale", float(round(stats["taxable_total"] + stats["vat_total"], 2))),
                         ("Da verificare", len(stats["anomalies"]))):
        put(ws, [label, value])
    put(ws, [])
    put(ws, ["Categoria", "Fatture", "Imponibile", "IVA", "Nota sulla deducibilità (indicativa)"], header=True)
    by_cat: dict[str, list] = defaultdict(lambda: [0, 0.0, 0.0])
    for r in in_month:
        c = by_cat[r["category"]]
        c[0] += 1
        c[1] += r["invoice"]["taxable"]
        c[2] += r["invoice"]["tax"]
    for k, v in sorted(by_cat.items(), key=lambda kv: -kv[1][1]):
        put(ws, [CATEGORY_IT.get(k, k), v[0], float(round(v[1], 2)), float(round(v[2], 2)),
                 TAX_HINTS.get(k, TAX_HINTS["other"])[1]])
    put(ws, [])
    put(ws, [DISCLAIMER_IT])
    for col, width in zip("ABCDE", (26, 10, 14, 12, 90)):
        ws.column_dimensions[col].width = width

    inv_ws = wb.create_sheet("Fatture")
    put(inv_ws, ["Data", "Numero", "Tipo", "Fornitore", "Partita IVA", "Categoria", "Fonte categoria",
                 "Imponibile", "IVA", "Totale", "Scadenza", "Pagata", "Da verificare"], header=True)
    for r in in_month:
        inv = r["invoice"]
        put(inv_ws, [inv["date"], inv["number"], inv.get("doc_type", ""), inv["supplier"]["name"],
                     inv["supplier"]["vat_number"], CATEGORY_IT.get(r["category"], r["category"]),
                     _source_it(r["category_source"]), float(inv["taxable"]), float(inv["tax"]),
                     float(inv["total"]), _due(inv), "sì" if r["paid"] else "no", len(flagged[r["id"]])])
    for col, width in zip("ABCDEFGHIJKLM", (11, 14, 7, 32, 14, 22, 14, 12, 10, 12, 11, 7, 12)):
        inv_ws.column_dimensions[col].width = width
    inv_ws.freeze_panes = "A2"

    chk = wb.create_sheet("Da verificare")
    put(chk, ["Fornitore", "Numero", "Data", "Problema", "Dettaglio", "Regola usata"], header=True)
    for r in in_month:
        for a in flagged[r["id"]]:
            put(chk, [r["invoice"]["supplier"]["name"], r["invoice"]["number"], r["invoice"]["date"],
                      ANOMALY_IT.get(a["kind"], a["kind"]), _msg(a), a.get("rule_it", "")])
    for col, width in zip("ABCDEF", (28, 14, 11, 30, 60, 70)):
        chk.column_dimensions[col].width = width
    for row in chk.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
